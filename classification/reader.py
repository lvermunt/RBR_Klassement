# Copyright (c) 2026
"""Read race results from supported file formats."""

import polars as pl
from openpyxl import load_workbook


class ResultReader:
    """Read race results from different file formats."""

    def __init__(self, file_format: str) -> None:
        """Initialize the ResultReader with the specified file format."""
        self.file_format = file_format

    def read_results(self, file_path: str) -> pl.DataFrame:
        """Read race results from the given file."""
        if self.file_format == "excel":
            return self._read_excel_results(file_path)
        if self.file_format == "text":
            return self._read_text_results(file_path)
        msg = "Unsupported file format. Only 'excel' and 'text' are supported."
        raise ValueError(msg)

    def _read_excel_results(self, file_path: str) -> pl.DataFrame:
        """Read a race result from an Excel file while preserving row-oriented worksheet data."""
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        sheet = workbook.active
        rows: list[tuple[object, ...]] = []

        for row in sheet.iter_rows(values_only=True):
            if all(value is None for value in row):
                continue
            rows.append(tuple(value for value in row))

        if not rows:
            return pl.DataFrame()

        width = max(len(row) for row in rows)
        padded_rows = [tuple(row + (None,) * (width - len(row))) for row in rows]
        return pl.DataFrame(padded_rows, orient="row")

    def _read_text_results(self, file_path: str) -> pl.DataFrame:
        """Read a race result from a tab-delimited text file."""
        return pl.read_csv(file_path, separator="\t")
